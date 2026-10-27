#!/usr/bin/env python3
"""
extract_launch_data.py — Extract Kumo Launch Tracker data from the SharePoint Excel
into a CSV file, then inject it into index.html as inline CSV_RAW.

Usage:
  python3 extract_launch_data.py

Prerequisites:
  pip install openpyxl

The script:
  1. Reads the local Excel file (downloaded via OneDrive)
  2. Parses the "LAUNCH MATRIX (NEW)" tab
  3. Writes data/launch_tracker.csv
  4. Injects the CSV into index.html as inline CSV_RAW (same pattern as kumo-engine-kpi)

Schedule: Run manually twice daily — 13:00 Europe/Skopje and 13:00 US/Pacific
"""

import sys, os, re, csv, io
from datetime import datetime

# Ensure openpyxl is importable
PYLIBS = os.path.expanduser("~/.personal-assistant/workspace/tmp/pylibs")
if os.path.isdir(PYLIBS):
    sys.path.insert(0, PYLIBS)

import openpyxl

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_FILE = os.path.join(SCRIPT_DIR, "data", "Kumo Team July-October Launch tracker.xlsx")
CSV_FILE   = os.path.join(SCRIPT_DIR, "data", "launch_tracker.csv")
HTML_FILE  = os.path.join(SCRIPT_DIR, "index.html")

SHEET_NAME = "LAUNCH MATRIX (NEW)"

def fmt_date(val):
    """Convert datetime or string to YYYY-MM-DD, or return empty."""
    if val is None:
        return ""
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d")
    s = str(val).strip()
    if s.lower() in ("n/a", "na", "none", ""):
        return ""
    # Try to parse common date strings
    try:
        return datetime.fromisoformat(s[:10]).strftime("%Y-%m-%d")
    except Exception:
        return ""

def extract_data(excel_path):
    """Parse the Excel and return a list of dicts."""
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb[SHEET_NAME]

    rows = []
    current_section = ""  # Month / milestone header from column A

    for r in range(5, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value   # Section header (month)
        b = ws.cell(row=r, column=2).value   # Sub-group
        c = ws.cell(row=r, column=3).value   # Product name
        d = ws.cell(row=r, column=4).value   # Version
        f = ws.cell(row=r, column=6).value   # PIC
        g = ws.cell(row=r, column=7).value   # Code Freeze Planned
        h = ws.cell(row=r, column=8).value   # Code Freeze Actual
        k = ws.cell(row=r, column=11).value  # Release Date Planned
        l = ws.cell(row=r, column=12).value  # Release Date Actual
        m = ws.cell(row=r, column=13).value  # Status

        # Update section when we see a header in column A
        if a and str(a).strip():
            header = str(a).strip()
            # Map section headers to month names
            if "august" in header.lower():
                current_section = "August"
            elif "september" in header.lower():
                current_section = "September"
            elif "oct" in header.lower():
                current_section = "October"
            elif "november" in header.lower():
                current_section = "November"
            elif "december" in header.lower():
                current_section = "December"
            elif "january" in header.lower():
                current_section = "January"
            elif "february" in header.lower():
                current_section = "February"
            elif "march" in header.lower():
                current_section = "March+"
            elif header.lower() in ("mile", "launch matrix"):
                continue
            else:
                current_section = header

        # Only process rows with a product name in column C
        if not c or not str(c).strip():
            continue

        product = str(c).strip().split("\n")[0].strip()  # Take first line
        version = str(d).strip().split("\n")[0].strip() if d else""
        pic = str(f).strip() if f else ""
        status = str(m).strip() if m else ""

        # Skip header row
        if product == "Model/Library":
            continue
        # Clean up product names
        product = product.replace('"', '')

        rows.append({
            "month": current_section,
            "product": product,
            "version": version,
            "pic": pic,
            "freeze_planned": fmt_date(g),
            "freeze_actual": fmt_date(h),
            "release_planned": fmt_date(k),
            "release_actual": fmt_date(l),
            "status": status,
        })

    return rows

def write_csv(rows, csv_path):
    """Write rows to CSV."""
    os.makedirs(os.path.dirname(csv_path), exist_ok=True)
    fieldnames = ["month","product","version","pic","freeze_planned","freeze_actual","release_planned","release_actual","status"]
    with open(csv_path, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)
    print(f"✅ Wrote {len(rows)} products to {csv_path}")

def inject_csv_into_html(csv_path, html_path):
    """Replace the CSV_RAW constant in index.html with the latest CSV data."""
    with open(csv_path, "r") as f:
        csv_content = f.read().strip()

    with open(html_path, "r") as f:
        html = f.read()

    # Replace the CSV_RAW block
    pattern = r"const CSV_RAW = `[^`]*`;"
    replacement = f"const CSV_RAW = `{csv_content}`;"

    if re.search(pattern, html):
        html = re.sub(pattern, replacement, html)
        print("✅ Injected CSV into index.html (replaced existing CSV_RAW)")
    else:
        print("⚠️  Could not find CSV_RAW in index.html — manual injection needed")
        return

    # Also update the last-updated timestamp
    now = datetime.now().strftime("%B %d, %Y %H:%M")
    ts_pattern = r"Last updated: [^·<]*"
    html = re.sub(ts_pattern, f"Last updated: {now} ", html)

    with open(html_path, "w") as f:
        f.write(html)
    print(f"✅ Updated timestamp to {now}")

def main():
    if not os.path.exists(EXCEL_FILE):
        print(f"❌ Excel file not found: {EXCEL_FILE}")
        print(f"   Download it first from SharePoint/OneDrive into data/")
        sys.exit(1)

    print(f"📊 Extracting from: {EXCEL_FILE}")
    rows = extract_data(EXCEL_FILE)
    write_csv(rows, CSV_FILE)
    inject_csv_into_html(CSV_FILE, HTML_FILE)
    print(f"\n🎉 Done! {len(rows)} products extracted and injected.")

if __name__ == "__main__":
    main()
